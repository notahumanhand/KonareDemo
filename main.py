from pathlib import Path

from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles


app = FastAPI()


# --------------------------------------------------
# Paths
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DEMO_DATA_DIR = BASE_DIR / "demo-data"
STATIC_DIR = BASE_DIR / "static"


# --------------------------------------------------
# Static files
# --------------------------------------------------

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


# --------------------------------------------------
# Demo case documents
# --------------------------------------------------

DOCUMENTS = [
    {
        "id": "01",
        "name": "Client Intake Questionnaire",
        "category": "Intake",
        "filename": "01_Client_Intake_Questionnaire.docx",
    },
    {
        "id": "02",
        "name": "Mock Passport Bio Page",
        "category": "Identity",
        "filename": "02_Mock_Passport_Bio_Page.docx",
    },
    {
        "id": "03",
        "name": "Mock Arrival / Admission Record",
        "category": "Immigration Record",
        "filename": "03_Mock_Arrival_Record.docx",
    },
    {
        "id": "04",
        "name": "Synthetic Birth Record",
        "category": "Civil Record",
        "filename": "04_Synthetic_Birth_Record.docx",
    },
    {
        "id": "05",
        "name": "Synthetic Marriage Record",
        "category": "Civil Record",
        "filename": "05_Synthetic_Marriage_Record.docx",
    },
    {
        "id": "06",
        "name": "Mock Spouse Passport Record",
        "category": "Identity",
        "filename": "06_Mock_Spouse_Passport_Record.docx",
    },
    {
        "id": "07",
        "name": "Synthetic Pay Statement",
        "category": "Financial Record",
        "filename": "07_Synthetic_Pay_Statement.docx",
    },
    {
        "id": "08",
        "name": "Synthetic Tax Summary",
        "category": "Financial Record",
        "filename": "08_Synthetic_Tax_Summary.docx",
    },
    {
        "id": "09",
        "name": "Synthetic Residential Lease Summary",
        "category": "Address Record",
        "filename": "09_Synthetic_Lease_Summary.docx",
    },
    {
        "id": "10",
        "name": "Synthetic Internal Case Note",
        "category": "Case Note",
        "filename": "10_Synthetic_Internal_Case_Note.docx",
    },
]


# --------------------------------------------------
# Helper functions
# --------------------------------------------------

def get_document(document_id: str):
    for document in DOCUMENTS:
        if document["id"] == document_id:
            return document

    return None


def extract_document_content(document):
    file_path = DEMO_DATA_DIR / document["filename"]

    if not file_path.is_file():
        raise HTTPException(
            status_code=500,
            detail="The requested demonstration document could not be found.",
        )

    docx = Document(file_path)

    paragraphs = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    fields = []

    for table in docx.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]

            if len(cells) >= 2 and cells[0]:
                fields.append(
                    {
                        "label": cells[0],
                        "value": cells[1],
                    }
                )

    return {
        "id": document["id"],
        "name": document["name"],
        "category": document["category"],
        "synthetic": True,
        "paragraphs": paragraphs,
        "fields": fields,
    }


# --------------------------------------------------
# Pages
# --------------------------------------------------

@app.get("/")
async def read_index():
    return FileResponse(STATIC_DIR / "index.html")


# --------------------------------------------------
# Document API
# --------------------------------------------------

@app.get("/api/cases/MG-2026-001/documents")
async def get_documents():
    return [
        {
            "id": document["id"],
            "name": document["name"],
            "category": document["category"],
        }
        for document in DOCUMENTS
    ]


@app.get("/api/cases/MG-2026-001/documents/{document_id}")
async def get_document_content(document_id: str):
    document = get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    return extract_document_content(document)